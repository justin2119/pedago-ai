"""Unit tests for the DOCX export of a LessonPlan."""

from io import BytesIO

from docx import Document

from pedago_ai.schemas.lesson_plan import (
    CANONICAL_MOMENTS,
    Activity,
    CapaciteContenu,
    DidacticMoment,
    EvaluationItem,
    LessonPlan,
    Objective,
    TeacherInfo,
    TimingSegment,
)
from pedago_ai.services.export_docx import lesson_plan_to_docx

ZIP_MAGIC = b"PK\x03\x04"


def _teacher():
    return TeacherInfo(
        nom="Dupont",
        prenom="Marie",
        dre="DRE Maritime",
        iesg="IESG Vogan",
        etablissement="Lycee de Vogan",
        telephone="90000000",
        classe="Seconde C",
        serie="C",
    )


def _full_plan():
    return LessonPlan(
        title="Le mouvement rectiligne uniforme",
        teacher=_teacher(),
        dre="DRE Maritime",
        etablissement="Lycee de Vogan",
        fiche_numero=3,
        semaine="Semaine 4",
        annee_scolaire="2026-2027",
        subject="Physique-Chimie",
        level="Seconde C",
        theme="Mecanique",
        sub_theme="Cinematique",
        competence="Resoudre un probleme de mouvement",
        duration_minutes=110,
        objectives=[
            Objective(description="Calculer une vitesse moyenne", domain="cognitif"),
            Objective(description="Representer un graphe x(t)", domain="psychomoteur"),
        ],
        prerequisites=["Notion de vitesse", "Proportionnalite"],
        materials=["Chronometre", "Regle graduee"],
        capacites_contenus=[
            CapaciteContenu(capacite="Calculer une vitesse", contenus=["Distance", "Duree"]),
            CapaciteContenu(capacite="Tracer un graphe", contenus=["Axes", "", "Echelle"]),
        ],
        timing=[TimingSegment(phase=m, duration_minutes=10) for m in CANONICAL_MOMENTS],
        activities=[
            Activity(
                phase=m,
                professor_activity=f"Activite du professeur ({m.value})",
                student_activity=f"Activite des eleves ({m.value})",
                duration_minutes=10,
            )
            for m in CANONICAL_MOMENTS
        ],
        evaluation=[
            EvaluationItem(question="Quelle est la formule de v ?", expected_answer="v = d / t"),
            EvaluationItem(question="Quelle unite pour v ?", expected_answer="m/s"),
        ],
        essentiel_a_retenir="La vitesse moyenne est le rapport distance sur duree.",
        exercices_application=[
            {"enonce": "Un mobile parcourt 100 m en 20 s. Calculer sa vitesse.", "correction": "v = 5 m/s"},
            {"enonce": "Exercice libre sans correction"},
        ],
        devoir_maison="Refaire l'exercice 2 et tracer le graphe x(t).",
        notes="Prevoir une calculatrice.",
    )


def _paragraph_texts(document):
    return [p.text for p in document.paragraphs]


def _all_cell_texts(document):
    return [cell.text for table in document.tables for row in table.rows for cell in row.cells]


def _find_table(document, first_header):
    for table in document.tables:
        if table.rows and table.rows[0].cells[0].text == first_header:
            return table
    raise AssertionError(f"Table with header {first_header!r} not found")


def test_full_plan_produces_non_empty_valid_docx_bytes():
    data = lesson_plan_to_docx(_full_plan())
    assert isinstance(data, bytes)
    assert len(data) > 0
    assert data.startswith(ZIP_MAGIC)
    # Must be re-openable as a valid Word document.
    document = Document(BytesIO(data))
    assert len(document.paragraphs) > 0


def test_full_plan_docx_can_be_written_and_reloaded_from_disk(tmp_path):
    path = tmp_path / "fiche.docx"
    path.write_bytes(lesson_plan_to_docx(_full_plan()))
    assert path.stat().st_size > 0
    reloaded = Document(str(path))
    assert "FICHE DE COURS" in _paragraph_texts(reloaded)


def test_full_plan_contains_all_sections():
    document = Document(BytesIO(lesson_plan_to_docx(_full_plan())))
    paragraphs = _paragraph_texts(document)
    for heading in [
        "Objectifs pedagogiques operationnels",
        "Prerequis",
        "Materiel",
        "Matrice capacites / contenus",
        "Deroulement",
        "Evaluation",
        "Essentiel a retenir",
        "Exercices d'application",
        "Devoir maison",
        "Notes",
    ]:
        assert heading in paragraphs, f"missing section {heading}"


def test_full_plan_contains_matrix_and_consolidation_content():
    document = Document(BytesIO(lesson_plan_to_docx(_full_plan())))
    cells = _all_cell_texts(document)
    assert "Calculer une vitesse" in cells
    assert "Distance\nDuree" in cells
    paragraphs = _paragraph_texts(document)
    assert "La vitesse moyenne est le rapport distance sur duree." in paragraphs
    assert "Refaire l'exercice 2 et tracer le graphe x(t)." in paragraphs
    assert "Correction : v = 5 m/s" in paragraphs
    assert "Exercice libre sans correction" in paragraphs


def test_deroulement_table_has_one_row_per_moment_in_canonical_order():
    document = Document(BytesIO(lesson_plan_to_docx(_full_plan())))
    table = _find_table(document, "Phase")
    assert len(table.rows) == 1 + len(CANONICAL_MOMENTS)
    phases = [row.cells[0].text for row in table.rows[1:]]
    assert phases == [m.value for m in CANONICAL_MOMENTS]
    assert table.rows[1].cells[1].text == "10 min"
    assert table.rows[1].cells[2].text == "Activite du professeur (prerequis)"


def test_header_table_uses_teacher_identity_and_lesson_info():
    document = Document(BytesIO(lesson_plan_to_docx(_full_plan())))
    header = _find_table(document, "DRE")
    values = {row.cells[0].text: row.cells[1].text for row in header.rows}
    assert values["DRE"] == "DRE Maritime"
    assert values["Nom et prenom"] == "Dupont Marie"
    assert values["Classe"] == "Seconde C - Serie C"
    assert values["Theme"] == "Mecanique - Cinematique"
    assert values["Duree"] == "110 min"


def test_plan_without_teacher_falls_back_to_plan_header_fields():
    plan = _full_plan()
    plan.teacher = None
    document = Document(BytesIO(lesson_plan_to_docx(plan)))
    header = _find_table(document, "DRE")
    values = {row.cells[0].text: row.cells[1].text for row in header.rows}
    assert values["DRE"] == "DRE Maritime"
    assert values["Etablissement"] == "Lycee de Vogan"
    assert values["Classe"] == "Seconde C"


def test_minimal_plan_without_optional_sections_still_exports():
    plan = LessonPlan(
        title="Minimal",
        subject="Physique-Chimie",
        level="Seconde C",
        theme="Mecanique",
        competence="Competence",
        duration_minutes=55,
        objectives=[Objective(description="Objectif", domain="cognitif")],
        timing=[TimingSegment(phase=DidacticMoment.PREREQUIS, duration_minutes=55)],
        activities=[Activity(phase=DidacticMoment.PREREQUIS, duration_minutes=55)],
    )
    data = lesson_plan_to_docx(plan)
    assert data.startswith(ZIP_MAGIC)
    document = Document(BytesIO(data))
    paragraphs = _paragraph_texts(document)
    assert "Matrice capacites / contenus" not in paragraphs
    assert "Devoir maison" not in paragraphs
    assert "Notes" not in paragraphs
