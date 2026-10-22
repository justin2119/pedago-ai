"""Ingestion pipeline for the pedagogical corpus.

Scans corpus/guides_execution/ and corpus/fiches_modeles/ (or the file and directory paths
given on the command line), extracts text from PDF, DOCX, TXT and MD files, cleans it,
segments it on pedagogical section markers, embeds each chunk through the OpenAI-compatible
proxy and upserts the rows into the Supabase document_chunks table.

Usage:
    uv run python -m pedago_ai.rag.ingest
    uv run python -m pedago_ai.rag.ingest corpus/guides_execution/lycee --dry-run
    uv run pedago-ai-ingest path/to/file.pdf
"""

from __future__ import annotations

import argparse
import hashlib
import logging
import re
import sys
import unicodedata
from pathlib import Path
from typing import Any

import httpx
from docx import Document
from pypdf import PdfReader
from tqdm import tqdm

from pedago_ai.config import Settings, get_settings
from pedago_ai.rag.chunking import chunk_text
from pedago_ai.rag.supabase_client import get_supabase

logger = logging.getLogger(__name__)

REPO_ROOT = Path(__file__).resolve().parents[3]
CORPUS_DIR = REPO_ROOT / "corpus"
DEFAULT_PATHS = [CORPUS_DIR / "guides_execution", CORPUS_DIR / "fiches_modeles"]
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}

CHUNK_MAX_CHARS = 1200
CHUNK_OVERLAP = 150
DEFAULT_BATCH_SIZE = 64
UPSERT_BATCH_SIZE = 100

TYPE_DIRS = {"guides_execution": "guide_execution", "fiches_modeles": "fiche_modele"}
TYPE_PREFIXES = {"guide": "guide_execution", "fiche_modele": "fiche_modele"}
MATIERE_ALIASES = {"sp": "sciences_physiques", "sciences_physiques": "sciences_physiques"}
CLASSE_ALIASES = {
    "seconde": "seconde",
    "2nde": "seconde",
    "2de": "seconde",
    "premiere": "premiere",
    "1ere": "premiere",
    "1re": "premiere",
    "terminale": "terminale",
    "tle": "terminale",
}
SERIE_CODES = {"a4", "s", "cd", "general"}
NIVEAUX_LYCEE = {"seconde", "premiere", "terminale"}

# Pedagogical section markers (Compétences, OPO, Situations/Activités d'apprentissage,
# Évaluation, ...) plus Markdown headings. A marker only counts when it starts a short line.
_SECTION_KEYWORDS = (
    r"comp[ée]tences?",
    r"opo\b",
    r"objectifs?\s+(?:d['’]apprentissage|p[ée]dagogiques)",
    r"situations?\s+d['’]apprentissage",
    r"activit[ée]s?\s+d['’]apprentissage",
    r"[ée]valuation",
    r"d[ée]roulement",
    r"contenus?\s+d['’]enseignement",
)
HEADING_RE = re.compile(
    r"^[ \t]*(?:#{1,6}[ \t]+.+|(?:\d+(?:\.\d+)*[.)]?[ \t]+)?(?:"
    + "|".join(_SECTION_KEYWORDS)
    + r").{0,120})[ \t]*$",
    re.IGNORECASE | re.MULTILINE,
)


def slugify(value: str) -> str:
    """Lowercase ASCII slug with underscores (accents removed)."""
    decomposed = unicodedata.normalize("NFKD", value)
    ascii_only = "".join(c for c in decomposed if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "_", ascii_only.lower()).strip("_")


def parse_metadata(path: Path) -> dict[str, Any]:
    """Derive document metadata from the folder layout first, then from the filename.

    Expected layout: corpus/<guides_execution|fiches_modeles>/lycee/<matiere>/<classe>/<serie>/
    Expected filename: <type>_<matiere>_<niveau>_<serie>[_<chapitre>].ext
    """
    parent_parts = [slugify(p) for p in path.resolve().parent.parts]
    dir_type: str | None = None
    tail: list[str] = []
    for index, part in enumerate(parent_parts):
        if part in TYPE_DIRS:
            dir_type = TYPE_DIRS[part]
            tail = parent_parts[index + 1 :]
            break

    niveau: str | None = None
    matiere: str | None = None
    classe: str | None = None
    serie: str | None = None
    for value in tail:
        if value == "lycee":
            niveau = "lycee"
        elif value in MATIERE_ALIASES and matiere is None:
            matiere = MATIERE_ALIASES[value]
        elif value in CLASSE_ALIASES and classe is None:
            classe = CLASSE_ALIASES[value]
        elif value in SERIE_CODES and serie is None:
            serie = value

    stem = slugify(path.stem)
    file_type: str | None = None
    for prefix, doc_type in TYPE_PREFIXES.items():
        if stem.startswith(prefix + "_"):
            file_type = doc_type
            stem = stem[len(prefix) + 1 :]
            break
    tokens = [t for t in stem.split("_") if t]

    if matiere is None:
        for alias in sorted(MATIERE_ALIASES, key=len, reverse=True):
            alias_tokens = alias.split("_")
            if tokens[: len(alias_tokens)] == alias_tokens:
                matiere = MATIERE_ALIASES[alias]
                tokens = tokens[len(alias_tokens) :]
                break
    if classe is None and tokens and tokens[0] in CLASSE_ALIASES:
        classe = CLASSE_ALIASES[tokens.pop(0)]
    if serie is None and tokens and tokens[0] in SERIE_CODES:
        serie = tokens.pop(0)
    if niveau is None and classe in NIVEAUX_LYCEE:
        niveau = "lycee"

    return {
        "type_document": dir_type or file_type or "document",
        "niveau": niveau,
        "classe": classe,
        "serie": serie or "non_applicable",
        "matiere": matiere,
        "chapitre": "_".join(tokens) or None,
        "titre": path.stem,
        "format": path.suffix.lower().lstrip("."),
    }


def extract_text(path: Path) -> str:
    """Return the raw text of a PDF, DOCX, TXT or MD file."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                parts.append(" | ".join(dict.fromkeys(c for c in cells if c)))
        return "\n".join(parts)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {path.name}")


def clean_text(raw: str) -> str:
    """Normalise whitespace, repair hyphenation and PDF line breaks."""
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\xa0", " ").replace("\u00ad", "")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    # A line break after a non-terminal word, followed by a lowercase start, is a PDF wrap.
    text = re.sub(r"(?<![.!?:;])\n(?=[a-zà-ÿ(])", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def segment_sections(text: str) -> list[tuple[str, str]]:
    """Split text on pedagogical headings, returning (section title, section body) pairs."""
    matches = list(HEADING_RE.finditer(text))
    if not matches:
        return [("document", text.strip())]

    sections: list[tuple[str, str]] = []
    preamble = text[: matches[0].start()].strip()
    if preamble:
        sections.append(("introduction", preamble))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        title = match.group(0).strip().lstrip("#").strip()[:120]
        sections.append((title, text[match.start() : end].strip()))
    return [(title, body) for title, body in sections if body]


def chunk_id(source: str, index: int, content: str) -> str:
    """Deterministic id so re-ingesting the same content upserts instead of duplicating."""
    return hashlib.sha256(f"{source}\x1f{index}\x1f{content}".encode()).hexdigest()


def build_chunks(text: str, metadata: dict[str, Any]) -> list[dict[str, Any]]:
    """Chunk each section on paragraph boundaries and attach metadata to every chunk."""
    pieces: list[tuple[str, str]] = []
    for title, body in segment_sections(text):
        for piece in chunk_text(body, max_chars=CHUNK_MAX_CHARS, overlap=CHUNK_OVERLAP):
            pieces.append((title, piece))

    source = metadata["source"]
    total = len(pieces)
    return [
        {
            "id": chunk_id(source, index, content),
            "content": content,
            "metadata": {
                **metadata,
                "section": title,
                "chunk_index": index,
                "chunk_count": total,
            },
        }
        for index, (title, content) in enumerate(pieces)
    ]


def embed_texts(http: httpx.Client, texts: list[str], settings: Settings) -> list[list[float]]:
    """Embed a batch of texts through the OpenAI-compatible proxy."""
    response = http.post(
        f"{settings.model_base_url.rstrip('/')}/embeddings",
        headers={"Authorization": f"Bearer {settings.model_api_key}"},
        json={"model": settings.embedding_model_name, "input": texts},
    )
    response.raise_for_status()
    items = sorted(response.json()["data"], key=lambda item: item["index"])
    vectors = [list(item["embedding"]) for item in items]
    if len(vectors) != len(texts):
        raise ValueError(f"Expected {len(texts)} embeddings, got {len(vectors)}")
    for vector in vectors:
        if len(vector) != settings.embedding_dimensions:
            raise ValueError(
                f"Embedding has {len(vector)} dimensions, expected {settings.embedding_dimensions}"
            )
    return vectors


def display_path(path: Path) -> str:
    """Repository-relative path when possible, used as the document source key."""
    resolved = path.resolve()
    try:
        return str(resolved.relative_to(REPO_ROOT))
    except ValueError:
        return str(resolved)


def discover_files(paths: list[Path]) -> list[Path]:
    """Expand directories into supported documents, skipping README files."""
    found: set[Path] = set()
    for path in paths:
        if path.is_dir():
            candidates = [p for p in path.rglob("*") if p.is_file()]
        elif path.is_file():
            candidates = [path]
        else:
            raise FileNotFoundError(f"No such file or directory: {path}")
        for candidate in candidates:
            if candidate.suffix.lower() in SUPPORTED_SUFFIXES and not candidate.name.startswith(
                "README"
            ):
                found.add(candidate.resolve())
    return sorted(found)


def ingest_file(
    path: Path,
    *,
    http: httpx.Client,
    supabase: Any,
    settings: Settings,
    batch_size: int,
    dry_run: bool,
) -> int:
    """Ingest one document and return the number of chunks produced."""
    source = display_path(path)
    metadata = {**parse_metadata(path), "source": source}
    text = clean_text(extract_text(path))
    if not text:
        logger.warning("No text extracted from %s; skipping", source)
        return 0

    records = build_chunks(text, metadata)
    if dry_run:
        logger.info("[dry-run] %s -> %d chunks, metadata=%s", source, len(records), metadata)
        return len(records)

    for start in range(0, len(records), batch_size):
        batch = records[start : start + batch_size]
        vectors = embed_texts(http, [record["content"] for record in batch], settings)
        for record, vector in zip(batch, vectors, strict=True):
            record["embedding"] = vector

    # Replace the document's previous chunks only after every embedding succeeded.
    table = settings.supabase_vector_table
    supabase.table(table).delete().eq("metadata->>source", source).execute()
    for start in range(0, len(records), UPSERT_BATCH_SIZE):
        supabase.table(table).upsert(
            records[start : start + UPSERT_BATCH_SIZE], on_conflict="id"
        ).execute()
    logger.info("Ingested %s -> %d chunks", source, len(records))
    return len(records)


def main(argv: list[str] | None = None) -> int:
    """CLI entry point. Returns a process exit code."""
    parser = argparse.ArgumentParser(description="Ingest the pedagogical corpus into Supabase.")
    parser.add_argument(
        "paths",
        nargs="*",
        type=Path,
        help="Files or directories to ingest (default: corpus guides and model sheets).",
    )
    parser.add_argument(
        "--dry-run", action="store_true", help="Parse and chunk only; no embeddings or writes."
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=DEFAULT_BATCH_SIZE,
        help="Chunks per embedding request.",
    )
    parser.add_argument("--log-level", default="INFO", help="Logging level.")
    args = parser.parse_args(argv)

    logging.basicConfig(
        level=args.log_level.upper(),
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
        stream=sys.stderr,
    )

    try:
        files = discover_files(args.paths or DEFAULT_PATHS)
    except FileNotFoundError as exc:
        logger.error("%s", exc)
        return 2
    if not files:
        logger.error("No supported documents (.pdf, .docx, .txt, .md) found.")
        return 1

    settings = get_settings()
    if not args.dry_run:
        missing = [
            name
            for name, value in (
                ("SUPABASE_URL", settings.supabase_url),
                ("SUPABASE_KEY", settings.supabase_key),
                ("MODEL_BASE_URL", settings.model_base_url),
                ("MODEL_API_KEY", settings.model_api_key),
            )
            if not value
        ]
        if missing:
            logger.error("Missing required settings: %s", ", ".join(missing))
            return 2

    supabase = None if args.dry_run else get_supabase()
    total_chunks = 0
    failures: list[str] = []
    with httpx.Client(timeout=60.0) as http:
        for path in tqdm(files, desc="Ingesting", unit="doc"):
            try:
                total_chunks += ingest_file(
                    path,
                    http=http,
                    supabase=supabase,
                    settings=settings,
                    batch_size=max(1, args.batch_size),
                    dry_run=args.dry_run,
                )
            except Exception:  # noqa: BLE001 - report and continue with the next document
                logger.exception("Failed to ingest %s", display_path(path))
                failures.append(display_path(path))

    logger.info(
        "Done: %d documents, %d chunks, %d failures", len(files), total_chunks, len(failures)
    )
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
