"""Export a LessonPlan to a formatted Word document."""

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from pedago_ai.schemas.lesson_plan import LessonPlan, TeacherInfo


def _add_header_table(document: Document, plan: LessonPlan) -> None:
    """Add the official administrative header table (teacher identity and lesson identification)."""
    teacher: TeacherInfo | None = plan.teacher
    rows = [
        ("DRE", teacher.dre if teacher else plan.dre),
        ("IESG", teacher.iesg if teacher else ""),
        ("Etablissement", teacher.etablissement if teacher else plan.etablissement),
        ("Nom et prenom", f"{teacher.nom} {teacher.prenom}" if teacher else ""),
        ("Contact / Numero", teacher.telephone if teacher else ""),
        ("Classe", f"{teacher.classe}" + (f" - Serie {teacher.serie}" if teacher and teacher.serie else "") if teacher else plan.level),
        ("Discipline / Matiere", plan.subject),
        ("Theme", plan.theme + (f" - {plan.sub_theme}" if plan.sub_theme else "")),
        ("Competence", plan.competence),
        ("Duree", f"{plan.duration_minutes} min"),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value or ""
        for run in cells[0].paragraphs[0].runs:
            run.bold = True


def _add_bullets(document: Document, items: list[str]) -> None:
    """Add one bullet paragraph per non-empty item."""
    for item in items:
        if str(item).strip():
            document.add_paragraph(str(item), style="List Bullet")


def _exercise_parts(item: Any) -> tuple[str, str]:
    """Return (enonce, correction) for an exercise item, accepting dict or plain text shapes."""
    if isinstance(item, dict):
        enonce = item.get("enonce") or item.get("question") or item.get("titre") or ""
        correction = item.get("correction") or item.get("reponse") or ""
        return str(enonce), str(correction)
    return str(item), ""


def lesson_plan_to_docx(plan: LessonPlan) -> bytes:
    """Render the lesson plan as a .docx file and return its bytes."""
    document = Document()
    republic = document.add_paragraph("REPUBLIQUE TOGOLAISE")
    republic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ministry = document.add_paragraph("Ministere des Enseignements Primaire, Secondaire, Technique et de la Formation Professionnelle")
    ministry.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("FICHE DE COURS", level=0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_header_table(document, plan)

    document.add_heading(plan.title, level=1)

    document.add_heading("Objectifs pedagogiques operationnels", level=1)
    for objective in plan.objectives:
        document.add_paragraph(objective.description, style="List Bullet")

    if plan.prerequisites:
        document.add_heading("Prerequis", level=1)
        _add_bullets(document, plan.prerequisites)

    if plan.materials:
        document.add_heading("Materiel", level=1)
        _add_bullets(document, plan.materials)

    if plan.capacites_contenus:
        document.add_heading("Matrice capacites / contenus", level=1)
        matrix = document.add_table(rows=1, cols=2)
        matrix.style = "Table Grid"
        for cell, header in zip(matrix.rows[0].cells, ["Capacite", "Contenus"], strict=True):
            cell.text = header
        for row_data in plan.capacites_contenus:
            row = matrix.add_row().cells
            row[0].text = row_data.capacite
            row[1].text = "\n".join(c for c in row_data.contenus if c.strip())

    document.add_heading("Deroulement", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Phase", "Duree", "Activite du professeur", "Activite des eleves"]
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        cell.text = header
    for activity in plan.activities:
        row = table.add_row().cells
        row[0].text = activity.phase.value
        row[1].text = f"{activity.duration_minutes} min"
        row[2].text = activity.professor_activity or ""
        row[3].text = activity.student_activity or ""

    if plan.evaluation:
        document.add_heading("Evaluation", level=1)
        for item in plan.evaluation:
            document.add_paragraph(f"{item.question} : {item.expected_answer}", style="List Number")

    if plan.essentiel_a_retenir.strip():
        document.add_heading("Essentiel a retenir", level=1)
        document.add_paragraph(plan.essentiel_a_retenir)

    if plan.exercices_application:
        document.add_heading("Exercices d'application", level=1)
        for item in plan.exercices_application:
            enonce, correction = _exercise_parts(item)
            if enonce.strip():
                document.add_paragraph(enonce, style="List Number")
            if correction.strip():
                document.add_paragraph(f"Correction : {correction}")

    if plan.devoir_maison.strip():
        document.add_heading("Devoir maison", level=1)
        document.add_paragraph(plan.devoir_maison)

    if plan.notes and plan.notes.strip():
        document.add_heading("Notes", level=1)
        document.add_paragraph(plan.notes)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
